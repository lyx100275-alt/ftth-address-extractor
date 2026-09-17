#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成「九级地址表」——地址树展开式 xlsx（每级节点各占一行 + 每户一行）

与 gen_addressbook.py 的区别：
  gen_addressbook.py  —— 每户一行（扁平表），列里直接写 9 级地址
  本脚本              —— 地址**树**展开：基础行 / 小区行 / 楼栋行 / 单元行 / 楼层行 / 户行
                         各级节点自身各占一行，叶子户行再各占一行。
                         这是资源/沙盘系统导入常用的形态，两者不可互替。

行数恒等式（自检用）：
    行数 = 1(基础) + 1(小区) + 楼栋数 + 单元数 + Σ层数 + Σ户数

表头与列序**必须**取自项目已定稿的表（--header-xlsx），不得自行重排。

输入 JSON 形态（任选）：
  {"地块名": {"1号楼": {"1": [18, 3], "2": [18, 3]}, "2号楼": {"1": [18, 2]}}}
  {"地块名": {"1号楼": {"1单元": {"层数": 18, "每层户数": 3}}}}

用法：
  python gen_9level_addressbook.py --site-json site.json \
      --header-xlsx "已定稿的标准地址.xlsx" --header-sheet <已定稿sheet名> \
      --addr "<分公司>,<一级>,<二级>,<三级>,<四级>" \
      --lv5-idx 5 --alias-idx <别名列号(0基，可选)> \
      --out 标准地址表.xlsx
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys


def load_header(xlsx, sheet, row=1):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    hdr = tuple(c if c is not None else '' for c in
                next(ws.iter_rows(min_row=row, max_row=row, values_only=True)))
    wb.close()
    return hdr


def normalize_bldg(b, units):
    """把「栋级标注」形态展开成逐单元配置。

    read_titleblock_households.py 的输出是**栋级**的：
        {"单元数": 2, "层数": 18, "每层户数": 3}
    图签上只有一条「N层/M户」标注，故当单元数为 n 时按同配置展开为 n 个单元，
    并明确告警：同一栋各单元层数/户数是否真的相同，图上没有逐单元证据，须逐单元核对。
    """
    if isinstance(units, dict) and {'层数', '每层户数'} <= set(units):
        n = int(units.get('单元数') or 1)
        cfg = [int(units['层数']), int(units['每层户数'])]
        if n > 1:
            print(f'  [告警] {bldg_label(b)}：栋级标注「{cfg[0]}层/{cfg[1]}户」按 {n} 个单元同配置展开，'
                  f'需逐单元核对（图上无逐单元层户证据）', file=sys.stderr)
        return {str(i): cfg for i in range(1, n + 1)}
    return units


def norm_units(u):
    """单元 -> (层数, 每层户数)；支持 [层,户] / {"层数":..,"每层户数":..}"""
    if isinstance(u, dict):
        if '层数' in u and '每层户数' in u:
            return int(u['层数']), int(u['每层户数'])
        raise ValueError(f'单元配置缺 层数/每层户数: {u}')
    if isinstance(u, (list, tuple)) and len(u) == 2:
        return int(u[0]), int(u[1])
    raise ValueError(f'无法识别的单元配置: {u!r}')


def unit_label(k):
    s = str(k).strip()
    m = re.search(r'(\d+)', s)
    return f'{m.group(1)}单元' if m else s


def bldg_label(k):
    s = str(k).strip()
    if re.search(r'号楼|#楼', s):
        return s
    m = re.search(r'(\d+)', s)
    return f'{m.group(1)}号楼' if m else s


def build_rows(site, blds, base5, header, alias=None, memo5='', lv5_idx=5, lv6_idx=7,
               lv7_idx=9, lv8_idx=11, lv9_idx=13, alias_idx=None, memo_idx=None,
               memo6=None):
    ncol = len(header)
    lv5 = site

    def mk(**kw):
        r = [''] * ncol
        for i, v in enumerate(base5):
            if i < ncol:
                r[i] = v
        r[lv5_idx] = kw.get('lv5', '')
        r[lv6_idx] = kw.get('lv6', '')
        r[lv7_idx] = kw.get('lv7', '')
        r[lv8_idx] = kw.get('lv8', '')
        r[lv9_idx] = kw.get('lv9', '')
        if alias_idx is not None:
            r[alias_idx] = kw.get('alias', '')
        if memo_idx is not None:
            r[memo_idx] = kw.get('memo', '')
        return r

    rows = [mk(), mk(lv5=lv5, alias=alias or '', memo=memo5)]
    blds = list(blds.items())
    for b, units in blds:
        rows.append(mk(lv5=lv5, lv6=bldg_label(b),
                       memo=(memo6(b) if memo6 else '')))
    for b, units in blds:
        for u in units:
            rows.append(mk(lv5=lv5, lv6=bldg_label(b), lv7=unit_label(u)))
    for b, units in blds:
        for u, cfg in units.items():
            f, _ = norm_units(cfg)
            for k in range(1, f + 1):
                rows.append(mk(lv5=lv5, lv6=bldg_label(b), lv7=unit_label(u),
                               lv8=f'{k}层'))
    for b, units in blds:
        for u, cfg in units.items():
            f, p = norm_units(cfg)
            for k in range(1, f + 1):
                for j in range(1, p + 1):
                    rows.append(mk(lv5=lv5, lv6=bldg_label(b), lv7=unit_label(u),
                                   lv8=f'{k}层', lv9=k * 100 + j))
    return rows


def main(argv=None):
    ap = argparse.ArgumentParser(description='生成九级地址表（地址树展开）')
    ap.add_argument('--site-json', required=True)
    ap.add_argument('--header-xlsx', required=True,
                    help='项目已定稿的标准地址表，取表头/列序（必需，禁止自造表头）')
    ap.add_argument('--header-sheet')
    ap.add_argument('--header-row', type=int, default=1)
    ap.add_argument('--addr', required=True,
                    help='前 5 级固定值，逗号分隔：分公司,一级,二级,三级,四级')
    ap.add_argument('--lv5-idx', type=int, default=5, help='五级列号(0基)')
    ap.add_argument('--lv6-idx', type=int, default=7)
    ap.add_argument('--lv7-idx', type=int, default=9)
    ap.add_argument('--lv8-idx', type=int, default=11)
    ap.add_argument('--lv9-idx', type=int, default=13)
    ap.add_argument('--alias-idx', type=int, default=None, help='别名列号(0基)')
    ap.add_argument('--memo-idx', type=int, default=None, help='备注列号(0基)')
    ap.add_argument('--lv6-memo-map', default=None,
                    help='楼栋->备注 映射 JSON 字符串，或 "前缀:逗号分隔楼栋" 简写')
    ap.add_argument('--out', required=True)
    ap.add_argument('--no-verify', action='store_true', help='跳过行数守恒自检')
    a = ap.parse_args(argv)

    base5 = tuple(x.strip() for x in a.addr.split(','))
    if len(base5) != 5:
        raise SystemExit('[参数错误] --addr 必须是 5 个逗号分隔值：分公司,一级,二级,三级,四级')

    with io.open(a.site_json, encoding='utf-8') as f:
        raw = json.load(f)
    sites = raw.get('地块', raw)

    memo6 = None
    if a.lv6_memo_map:
        txt = a.lv6_memo_map
        if txt.strip().startswith('{'):
            mm = json.loads(txt)
        else:
            pre, blds = txt.split(':', 1)
            mm = {b.strip(): pre.strip() for b in blds.split(',')}
        memo6 = (lambda b: mm.get(str(b), mm.get(re.sub(r'\D', '', str(b)), '')))

    header = load_header(a.header_xlsx, a.header_sheet, a.header_row)
    print(f'[表头] 取自 {os.path.basename(a.header_xlsx)} / {a.header_sheet or "首个sheet"}：'
          f'{len(header)} 列')

    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for site, blds in sites.items():
        blds = {b: normalize_bldg(b, u) for b, u in blds.items()}
        rows = build_rows(site, blds, base5, header, alias=raw.get('别名', {}).get(site) if isinstance(raw.get('别名'), dict) else None,
                          memo5='', lv5_idx=a.lv5_idx, lv6_idx=a.lv6_idx,
                          lv7_idx=a.lv7_idx, lv8_idx=a.lv8_idx, lv9_idx=a.lv9_idx,
                          alias_idx=a.alias_idx, memo_idx=a.memo_idx, memo6=memo6)
        # 守恒自检
        n_unit = sum(len(u) for u in blds.values())
        n_floor = sum(norm_units(c)[0] for u in blds.values() for c in u.values())
        n_hu = sum(norm_units(c)[0] * norm_units(c)[1]
                   for u in blds.values() for c in u.values())
        expect = 2 + len(blds) + n_unit + n_floor + n_hu
        if not a.no_verify and len(rows) != expect:
            raise SystemExit(f'[守恒校验失败] {site}: 实得 {len(rows)} 行 ≠ '
                             f'应有 {expect} 行（基础1+小区1+楼栋{len(blds)}+单元{n_unit}+'
                             f'楼层{n_floor}+户{n_hu}）')
        ws = wb.create_sheet(site[:31])
        ws.append(list(header))
        for r in rows:
            ws.append(r)
        idxs = [a.lv5_idx, a.lv6_idx, a.lv7_idx, a.lv8_idx, a.lv9_idx]
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, size=10)
            cell.fill = PatternFill('solid', fgColor='DDEBF7')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        for r in range(2, len(rows) + 2):
            depth = sum(1 for i in idxs if ws.cell(row=r, column=i + 1).value not in (None, ''))
            for c in range(1, len(header) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.font = Font(size=10, bold=(depth == 1))
                cell.alignment = Alignment(horizontal='left', vertical='center')
                if depth == 1:
                    cell.fill = PatternFill('solid', fgColor='F2F2F2')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(header))}{len(rows) + 1}'
        print(f'  {site}: 楼栋{len(blds)} 单元{n_unit} 楼层{n_floor} 户{n_hu} → 行{len(rows)} ✓')

    wb.save(a.out)
    print(f'[写出] {a.out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
